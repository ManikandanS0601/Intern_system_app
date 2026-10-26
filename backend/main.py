import os, re, io, json, uuid, hashlib, secrets, smtplib, time
from pathlib import Path
from datetime import datetime
from email.message import EmailMessage
from typing import Optional, List
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Header
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy import create_engine, Column, String, Integer, Float, Text, DateTime, Boolean
from sqlalchemy.orm import declarative_base, sessionmaker
from openpyxl import Workbook

DATA = os.getenv("DATA_DIR", "data"); os.makedirs(f"{DATA}/resumes", exist_ok=True)
engine = create_engine(os.getenv("DATABASE_URL", f"sqlite:///{DATA}/app.db"), connect_args={"check_same_thread": False})
Session = sessionmaker(engine); Base = declarative_base()
now = datetime.utcnow
MAX_ROUNDS = 3
THRESHOLD = float(os.getenv("PASS_THRESHOLD", 80))
XLSX = f"{DATA}/workbook.xlsx"

class User(Base):
    __tablename__ = "users"; id = Column(Integer, primary_key=True); email = Column(String, unique=True); pw = Column(String); role = Column(String)
class Token(Base):
    __tablename__ = "tokens"; token = Column(String, primary_key=True); user_id = Column(Integer)
class Intern(Base):
    __tablename__ = "interns"; id = Column(String, primary_key=True); user_id = Column(Integer); name = Column(String); email = Column(String, unique=True)
    phone = Column(String); education = Column(Text); skills = Column(Text); department = Column(String); role = Column(String)
    resume = Column(String); joined = Column(DateTime, default=now); decision = Column(String, default="ACTIVE")
class Task(Base):
    __tablename__ = "tasks"; id = Column(String, primary_key=True); intern_id = Column(String); title = Column(String); description = Column(Text)
    criteria = Column(Text); deadline = Column(String); round = Column(Integer, default=0); status = Column(String, default="ASSIGNED"); assigned = Column(DateTime, default=now)
class Submission(Base):
    __tablename__ = "submissions"; id = Column(Integer, primary_key=True); task_id = Column(String); round = Column(Integer); content = Column(Text); links = Column(Text); created = Column(DateTime, default=now)
class Review(Base):  # append-only: never updated except admin_decision
    __tablename__ = "reviews"; id = Column(Integer, primary_key=True); task_id = Column(String); intern_id = Column(String); submission_id = Column(Integer); round = Column(Integer)
    score = Column(Float); checklist = Column(Text); feedback = Column(Text); recommendation = Column(String); ai_mode = Column(String); admin_decision = Column(String); created = Column(DateTime, default=now)
class EmailLog(Base):
    __tablename__ = "emails"; id = Column(Integer, primary_key=True); intern_id = Column(String); recipient = Column(String); subject = Column(String)
    type = Column(String); sent = Column(DateTime, default=now); status = Column(String); task_id = Column(String); key = Column(String, unique=True)
class Feedback(Base):
    __tablename__ = "feedback"; id = Column(Integer, primary_key=True); intern_id = Column(String); task_id = Column(String); category = Column(String)
    ratings = Column(Text); description = Column(Text); challenges = Column(Text); suggestions = Column(Text); comments = Column(Text)
    confidential = Column(Boolean, default=False); created = Column(DateTime, default=now); status = Column(String, default="Open"); notes = Column(Text, default="")
class Audit(Base):
    __tablename__ = "audit"; id = Column(Integer, primary_key=True); actor = Column(String); action = Column(String); detail = Column(Text); created = Column(DateTime, default=now)

def hp(p, salt=None):
    salt = salt or secrets.token_hex(8); return salt + "$" + hashlib.pbkdf2_hmac("sha256", p.encode(), salt.encode(), 100000).hex()
def vp(p, h): return hp(p, h.split("$")[0]) == h

app = FastAPI(title="Intern Task Management & Evaluation")

@app.get("/")
def health():
    return {"status": "ok", "docs": "/docs"}

def db():
    s = Session()
    try: yield s
    finally: s.close()
def audit(s, actor, action, detail=""): s.add(Audit(actor=actor, action=action, detail=detail))
def user(authorization: str = Header(""), s=Depends(db)):
    t = s.get(Token, authorization.replace("Bearer ", ""))
    u = s.get(User, t.user_id) if t else None
    if not u: raise HTTPException(401, "Not authenticated")
    return u
def admin(u=Depends(user)):
    if u.role != "admin": raise HTTPException(403, "Admin only")
    return u
def me(u, s):
    i = s.query(Intern).filter_by(user_id=u.id).first()
    if not i: raise HTTPException(403, "Not an intern")
    return i

@app.on_event("startup")
def startup():
    Base.metadata.create_all(engine); s = Session()
    if not s.query(User).filter_by(role="admin").first():
        s.add(User(email=os.getenv("ADMIN_EMAIL", "admin@example.com"), pw=hp(os.getenv("ADMIN_PASSWORD", "change-me")), role="admin")); s.commit()
    sync_excel(s); s.close()

class Login(BaseModel): email: str; password: str
@app.post("/auth/login")
def login(b: Login, s=Depends(db)):
    u = s.query(User).filter_by(email=b.email).first()
    if not u or not vp(b.password, u.pw): raise HTTPException(401, "Bad credentials")
    t = secrets.token_urlsafe(24); s.add(Token(token=t, user_id=u.id)); s.commit(); return {"token": t, "role": u.role}

# ---------- Resume extraction ----------
SKILLS = ["java", "spring", "python", "fastapi", "react", "javascript", "sql", "mongodb", "docker", "html", "css", "machine learning", "pandas", "aws"]
def extract_text(name, data):
    if name.lower().endswith(".pdf"):
        from pypdf import PdfReader; return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if name.lower().endswith(".docx"):
        import docx; return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    raise HTTPException(415, "Only PDF or DOCX")
@app.post("/resumes/extract")
async def extract(f: UploadFile = File(...), _=Depends(admin)):
    data = await f.read()
    if len(data) > 5_000_000: raise HTTPException(413, "File too large")
    txt = extract_text(f.filename, data)
    em = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", txt); ph = re.search(r"\+?\d[\d\s-]{8,14}\d", txt)
    stored = f"{uuid.uuid4().hex}_{os.path.basename(f.filename)}"; open(f"{DATA}/resumes/{stored}", "wb").write(data)
    d = {"name": txt.strip().split("\n")[0][:80] if txt.strip() else None, "email": em.group(0) if em else None, "phone": ph.group(0) if ph else None,
         "skills": [k for k in SKILLS if k in txt.lower()], "resume": stored}
    d["missing_fields"] = [k for k in ("name", "email", "phone") if not d[k]]  # flagged, never invented
    return {"draft": d, "note": "Review and correct, then POST /interns to confirm."}

class InternIn(BaseModel): name: str; email: str; phone: str = ""; education: str = ""; skills: List[str] = []; department: str = ""; role: str = ""; resume: str = ""
@app.post("/interns", status_code=201)
def create_intern(b: InternIn, a=Depends(admin), s=Depends(db)):
    email = b.email
    if b.resume:
        if Path(b.resume).name != b.resume or "/" in b.resume or "\\" in b.resume:
            raise HTTPException(422, "Uploaded resume was not found; upload it again before saving the intern")
        resume_path = Path(DATA) / "resumes" / b.resume
        if not resume_path.is_file():
            raise HTTPException(422, "Uploaded resume was not found; upload it again before saving the intern")
        resume_text = extract_text(resume_path.name, resume_path.read_bytes())
        resume_email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", resume_text)
        if not resume_email:
            raise HTTPException(422, "No email address found in the uploaded resume")
        email = resume_email.group(0)
    if not re.fullmatch(r"[\w.+-]+@[\w-]+\.[\w.-]+", email): raise HTTPException(422, "Invalid email")
    email = email.lower()
    if s.query(Intern).filter_by(email=email).first(): raise HTTPException(409, "Duplicate intern")
    pw = secrets.token_urlsafe(8); u = User(email=email, pw=hp(pw), role="intern"); s.add(u); s.flush()
    i = Intern(id=f"INT-{s.query(Intern).count()+1:04d}", user_id=u.id, **{**b.dict(), "email": email, "skills": ", ".join(b.skills)}); s.add(i)
    audit(s, a.email, "create_intern", i.id); s.commit()
    email_status = send_email(s, i, "Welcome to the Internship Program", f"Hello {i.name},\n\nWelcome! Your internship portal account has been created.\n\nEmail: {i.email}\nTemporary password: {pw}\n\nPlease sign in to the internship portal and change your password if prompted.\n\nInternship Management Team", "WELCOME", None)
    sync_excel(s); return {"intern_id": i.id, "temporary_password": pw, "email_to": i.email, "email_status": email_status}

# ---------- Email (mock unless SMTP_HOST set; idempotent via key) ----------
def send_email(s, intern, subject, body, typ, task_id, rnd=0):
    key = f"{intern.id}|{typ}|{task_id}|{rnd}"
    if s.query(EmailLog).filter_by(key=key).first(): return "DUPLICATE_SKIPPED"
    host = os.getenv("SMTP_HOST"); status = "MOCK_NOT_SENT"
    if host:
        missing = [name for name in ("SMTP_FROM", "SMTP_USER", "SMTP_PASSWORD") if not os.getenv(name)]
        if missing:
            status = f"NOT_CONFIGURED: missing {', '.join(missing)}"
        else:
            try:
                m = EmailMessage(); m["Subject"], m["To"], m["From"] = subject, intern.email, os.getenv("SMTP_FROM"); m.set_content(body)
                with smtplib.SMTP(host, int(os.getenv("SMTP_PORT", 587)), timeout=15) as c:
                    c.starttls(); c.login(os.getenv("SMTP_USER"), os.getenv("SMTP_PASSWORD")); c.send_message(m)
                status = "SENT"
            except Exception as e: status = f"FAILED: {e}"[:200]
    if status.startswith("FAILED"): key += f"|{time.time()}"  # allow retry
    s.add(EmailLog(intern_id=intern.id, recipient=intern.email, subject=subject, type=typ, status=status, task_id=task_id, key=key)); s.commit(); return status

# ---------- Tasks ----------
class Crit(BaseModel): text: str; mandatory: bool = True
class TaskIn(BaseModel): intern_id: str; title: str; description: str; criteria: List[Crit]; deadline: str
@app.get("/interns")
def list_interns(q: str = "", _=Depends(admin), s=Depends(db)):
    return [dict(id=i.id, name=i.name, email=i.email, role=i.role, decision=i.decision, has_resume=bool(i.resume)) for i in s.query(Intern).all() if q.lower() in (i.name + i.skills).lower()]
@app.get("/interns/{iid}/resume")
def get_intern_resume(iid: str, u=Depends(user), s=Depends(db)):
    i = s.get(Intern, iid)
    if not i: raise HTTPException(404, "Intern not found")
    if u.role != "admin" and me(u, s).id != i.id: raise HTTPException(403, "Not allowed to access this resume")
    if not i.resume or Path(i.resume).name != i.resume or "/" in i.resume or "\\" in i.resume:
        raise HTTPException(404, "Resume not found")
    resume = (Path(DATA) / "resumes" / i.resume).resolve()
    if resume.parent != (Path(DATA) / "resumes").resolve() or not resume.is_file():
        raise HTTPException(404, "Resume not found")
    media_type = "application/pdf" if resume.suffix.lower() == ".pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return FileResponse(resume, media_type=media_type, filename=resume.name)
@app.get("/interns/{iid}/task-recommendations")
def recommend(iid: str, _=Depends(admin), s=Depends(db)):
    i = s.get(Intern, iid); sk = (i.skills + " " + i.role).lower()  # skills/role only, no personal attributes
    r = []
    if "java" in sk or "spring" in sk: r.append("Implement a REST API with validation and database persistence")
    if "react" in sk: r.append("Build a responsive dashboard with specified components")
    if "pandas" in sk or "sql" in sk: r.append("Clean a dataset, analyze trends and write a report")
    return {"ai_mode": "mock-heuristic", "suggestions": r or ["General coding exercise"], "requires_admin_approval": True}
@app.post("/tasks", status_code=201)
def create_task(b: TaskIn, a=Depends(admin), s=Depends(db)):
    i = s.get(Intern, b.intern_id)
    if not i: raise HTTPException(404, "Intern not found")
    t = Task(id=f"TSK-{s.query(Task).count()+1:04d}", intern_id=i.id, title=b.title, description=b.description, criteria=json.dumps([c.dict() for c in b.criteria]), deadline=b.deadline)
    s.add(t); audit(s, a.email, "assign_task", t.id); s.commit()
    reqs = "\n".join(f"{n}. {c.text}" for n, c in enumerate(b.criteria, 1))
    email_status = send_email(s, i, f"New Task Assigned – {t.title}", f"Hello {i.name},\n\nTask: {t.title}\nDescription: {t.description}\nRequirements:\n{reqs}\nDeadline: {t.deadline}\n\nSubmit via the internship portal. Up to {MAX_ROUNDS} review rounds apply.\n\nInternship Management Team", "TASK_ASSIGNED", t.id)
    sync_excel(s); return {"task_id": t.id, "email_status": email_status}
@app.get("/tasks")
def my_tasks(u=Depends(user), s=Depends(db)):
    q = s.query(Task) if u.role == "admin" else s.query(Task).filter_by(intern_id=me(u, s).id)
    return [dict(id=t.id, title=t.title, status=t.status, round=t.round, deadline=t.deadline, criteria=json.loads(t.criteria)) for t in q]

# ---------- AI evaluation (provider adapter) ----------
def ai_evaluate(task, text, links):
    """MOCK: keyword-evidence heuristic. Replace branch for AI_PROVIDER!=mock with an LLM call that
    receives criteria as trusted instructions and the submission as quoted, untrusted data."""
    if len((text + links).strip()) < 20: return None
    low = (text + " " + links).lower(); res = []
    for c in json.loads(task.criteria):
        w = re.findall(r"[a-z]{4,}", c["text"].lower()); hit = [x for x in w if x in low]
        res.append({"requirement": c["text"], "mandatory": c["mandatory"], "result": "PASSED" if w and len(hit) / len(w) >= .6 else "FAILED", "evidence": hit})
    return res
class Sub(BaseModel): content: str; links: str = ""
@app.post("/tasks/{tid}/submit")
def submit(tid: str, b: Sub, u=Depends(user), s=Depends(db)):
    t = s.get(Task, tid); i = me(u, s)
    if not t or t.intern_id != i.id: raise HTTPException(404, "Task not found")
    if t.status not in ("ASSIGNED", "NEEDS_REVISION"): raise HTTPException(409, f"Task is {t.status}")
    res = ai_evaluate(t, b.content, b.links)
    if res is None: return {"recommendation": "INSUFFICIENT_EVIDENCE", "message": "Please provide more evidence (files, repo/deploy links). No review round consumed."}
    rnd = t.round + 1; sub = Submission(task_id=t.id, round=rnd, content=b.content, links=b.links); s.add(sub); s.flush()
    score = 100 * sum(r["result"] == "PASSED" for r in res) / len(res)
    ok = all(r["result"] == "PASSED" for r in res if r["mandatory"]) and score >= THRESHOLD
    failed = [r["requirement"] for r in res if r["result"] != "PASSED"]
    t.round = rnd; t.status = "COMPLETED" if ok else ("REVIEW_FAILED" if rnd >= MAX_ROUNDS else "NEEDS_REVISION")
    fb = "All mandatory requirements met." if ok else "Unmet: " + "; ".join(failed)
    rev = Review(task_id=t.id, intern_id=i.id, submission_id=sub.id, round=rnd, score=score, checklist=json.dumps(res), feedback=fb,
                 recommendation="PASS" if ok else ("RECOMMEND_REJECTION_PENDING_ADMIN" if t.status == "REVIEW_FAILED" else "NEEDS_REVISION"), ai_mode=os.getenv("AI_PROVIDER", "mock")); s.add(rev); s.commit()
    typ = {"COMPLETED": "COMPLETION", "NEEDS_REVISION": "REVIEW_FEEDBACK", "REVIEW_FAILED": "REVIEW_FEEDBACK"}[t.status]
    send_email(s, i, f"Review result – {t.title} (Round {rnd})", fb, typ, t.id, rnd); sync_excel(s)
    return {"status": t.status, "round": rnd, "score": score, "checklist": res, "feedback": fb, "recommendation": rev.recommendation}
@app.get("/tasks/{tid}/reviews")
def reviews(tid: str, u=Depends(user), s=Depends(db)):
    t = s.get(Task, tid)
    if not t or (u.role != "admin" and t.intern_id != me(u, s).id): raise HTTPException(404, "Not found")
    return [dict(round=r.round, score=r.score, checklist=json.loads(r.checklist), feedback=r.feedback, recommendation=r.recommendation, admin_decision=r.admin_decision) for r in s.query(Review).filter_by(task_id=tid)]
class Decision(BaseModel): decision: str  # REJECT | RETAIN | EXTEND
@app.post("/interns/{iid}/decision")
def decide(iid: str, b: Decision, a=Depends(admin), s=Depends(db)):
    i = s.get(Intern, iid)
    if b.decision not in ("REJECT", "RETAIN", "EXTEND"): raise HTTPException(422, "Bad decision")
    if b.decision == "REJECT" and not s.query(Task).filter_by(intern_id=iid, status="REVIEW_FAILED").first(): raise HTTPException(409, "No failed review to act on")
    if b.decision == "EXTEND":  # exception process: grant one extra round
        t = s.query(Task).filter_by(intern_id=iid, status="REVIEW_FAILED").first(); t.status = "NEEDS_REVISION"; t.round = MAX_ROUNDS - 1
    else: i.decision = "REJECTED" if b.decision == "REJECT" else "ACTIVE"
    audit(s, a.email, "intern_decision", f"{iid}:{b.decision}"); s.commit()
    if b.decision == "REJECT": send_email(s, i, "Internship decision", "After administrator review, the internship will not continue.", "FINAL_REJECTION", None)
    sync_excel(s); return {"intern": iid, "decision": b.decision}

# ---------- Feedback (never read by evaluation code) ----------
CATS = ["Task Assignment", "Task Requirements and Clarity", "AI Evaluation and Review Process", "Mentor or Reviewer Support", "Email Notifications and Communication", "Workload and Deadlines", "Internship Learning Experience", "Technical Issues", "General Feedback"]
RK = ["overall", "difficulty", "clarity", "ai_fairness", "reviewer_support", "email"]
class FB(BaseModel):
    category: str; task_id: Optional[str] = None; ratings: dict; description: str = ""; challenges: str = ""; suggestions: str = ""; comments: str = ""; confidential: bool = False
@app.post("/feedback", status_code=201)
def post_fb(b: FB, u=Depends(user), s=Depends(db)):
    i = me(u, s)
    if b.category not in CATS or any(not 1 <= int(b.ratings.get(k, 3)) <= 5 for k in RK): raise HTTPException(422, "Invalid category/rating")
    if b.task_id and (not s.get(Task, b.task_id) or s.get(Task, b.task_id).intern_id != i.id): raise HTTPException(404, "Task not found")
    f = Feedback(intern_id=i.id, **b.dict()); f.ratings = json.dumps(b.ratings); s.add(f); s.commit()
    send_email(s, i, "Thanks for your feedback", "We received your feedback. Access: authorized administrators only.", "FEEDBACK_CONFIRMATION", b.task_id, f.id); sync_excel(s); return {"feedback_id": f.id}
def fb_out(f, s, full=False):
    i = s.get(Intern, f.intern_id); hide = f.confidential
    return dict(id=f.id, intern_id=None if hide else f.intern_id, intern_name=None if hide else i.name, task_id=f.task_id, category=f.category, ratings=json.loads(f.ratings),
                description=f.description, challenges=f.challenges, suggestions=f.suggestions, confidential=f.confidential, status=f.status, notes=f.notes, created=str(f.created))
@app.get("/feedback/mine")
def my_fb(u=Depends(user), s=Depends(db)): return [fb_out(f, s) for f in s.query(Feedback).filter_by(intern_id=me(u, s).id)]
@app.get("/admin/feedback")
def all_fb(category: str = "", a=Depends(admin), s=Depends(db)):
    audit(s, a.email, "view_feedback", category); s.commit()
    return [fb_out(f, s) for f in s.query(Feedback) if not category or f.category == category]
class FBUpd(BaseModel): status: Optional[str] = None; note: str = ""
@app.patch("/admin/feedback/{fid}")
def upd_fb(fid: int, b: FBUpd, a=Depends(admin), s=Depends(db)):
    f = s.get(Feedback, fid)
    if b.status:
        if b.status not in ("Open", "Under Review", "In Progress", "Resolved"): raise HTTPException(422, "Bad status")
        f.status = b.status
    if b.note: f.notes += f"\n[{now():%F %T} {a.email}] {b.note}"
    audit(s, a.email, "update_feedback", str(fid)); s.commit(); sync_excel(s); return fb_out(f, s)
@app.get("/admin/feedback/analysis")
def analysis(a=Depends(admin), s=Depends(db)):
    fs = s.query(Feedback).all(); n = len(fs); avg = {k: round(sum(json.loads(f.ratings).get(k, 0) for f in fs) / n, 2) for k in RK} if n else {}
    themes = {t: sum(t in (f.description + f.challenges + f.suggestions).lower() for f in fs) for t in ("unclear", "deadline", "email", "bug", "slow", "incomplete")}
    return {"AI_GENERATED": True, "mode": "mock keyword themes", "responses": n, "averages": avg, "theme_mentions": themes,
            "caveat": "Reported experiences, not verified facts. Verify against task/review records before acting. Small n may identify authors.",
            "suggestions": ["Review task descriptions/criteria"] if themes["unclear"] >= 2 else []}

# ---------- Dashboard & Excel ----------
@app.get("/dashboard")
def dash(_=Depends(admin), s=Depends(db)):
    T = s.query(Task).all(); st = lambda x: sum(t.status == x for t in T)
    return {"interns": s.query(Intern).count(), "tasks": len(T), "completed": st("COMPLETED"), "needs_revision": st("NEEDS_REVISION"), "attention_review_failed": st("REVIEW_FAILED"),
            "rounds": {r: sum(t.round == r for t in T) for r in range(4)}, "recent_emails": [dict(to=e.recipient, type=e.type, status=e.status) for e in s.query(EmailLog).order_by(EmailLog.id.desc()).limit(10)]}
def sync_excel(s):
    """DB is source of truth; workbook is rebuilt atomically (temp file + replace) with retry on lock."""
    wb = Workbook(); wb.remove(wb.active)
    def sheet(name, head, rows):
        w = wb.create_sheet(name); w.append(head); [w.append(r) for r in rows]
    sheet("Intern Records", ["Intern ID", "Name", "Email", "Phone", "Education", "Skills", "Department", "Role", "Resume Filename", "Joining Date", "Assigned Task ID", "Current Review Round", "Task Status", "Final Decision"],
          [[i.id, i.name, i.email, i.phone, i.education, i.skills, i.department, i.role, i.resume, str(i.joined), (t.id if t else ""), (t.round if t else 0), (t.status if t else ""), i.decision]
           for i in s.query(Intern) for t in [s.query(Task).filter_by(intern_id=i.id).order_by(Task.assigned.desc()).first()]])
    sheet("Task Records", ["Task ID", "Intern ID", "Title", "Description", "Requirements", "Assigned", "Deadline", "Round", "Status"],
          [[t.id, t.intern_id, t.title, t.description, "; ".join(c["text"] for c in json.loads(t.criteria)), str(t.assigned), t.deadline, t.round, t.status] for t in s.query(Task)])
    sheet("Review History", ["Review ID", "Intern ID", "Task ID", "Round", "Date", "AI Score", "Passed", "Failed", "Feedback", "AI Recommendation", "Admin Decision"],
          [[r.id, r.intern_id, r.task_id, r.round, str(r.created), r.score, "; ".join(c["requirement"] for c in json.loads(r.checklist) if c["result"] == "PASSED"),
            "; ".join(c["requirement"] for c in json.loads(r.checklist) if c["result"] != "PASSED"), r.feedback, r.recommendation, r.admin_decision] for r in s.query(Review)])
    sheet("Email History", ["Email ID", "Intern ID", "Recipient", "Subject", "Type", "Sent", "Delivery Status", "Task ID"],
          [[e.id, e.intern_id, e.recipient, e.subject, e.type, str(e.sent), e.status, e.task_id] for e in s.query(EmailLog)])
    sheet("Candidate Feedback", ["Feedback ID", "Intern ID", "Intern Name", "Task ID", "Category", "Overall", "Difficulty", "Clarity", "AI Fairness", "Reviewer Support", "Email", "Description", "Challenges", "Suggestions", "Comments", "Confidentiality", "Submitted", "Status"],
          [[f.id, o["intern_id"] or "CONFIDENTIAL", o["intern_name"] or "CONFIDENTIAL", f.task_id, f.category, *[o["ratings"].get(k) for k in RK], f.description, f.challenges, f.suggestions, f.comments,
            "Confidential" if f.confidential else "Standard", str(f.created), f.status] for f in s.query(Feedback) for o in [fb_out(f, s)]])
    for n in range(5):
        try: wb.save(XLSX + ".tmp"); os.replace(XLSX + ".tmp", XLSX); return
        except PermissionError: time.sleep(.3)
@app.get("/export/excel")
def export(a=Depends(admin), s=Depends(db)): sync_excel(s); audit(s, a.email, "export_excel"); s.commit(); return FileResponse(XLSX, filename="intern_records.xlsx")

import chat; chat.register(app, globals())
